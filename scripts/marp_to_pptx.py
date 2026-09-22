"""Build a PowerPoint from a Chicago Marp deck.

Source of truth is docs/slides/chicago-sept-24/dayN/dayN.md. Each Marp slide becomes
one PowerPoint slide: the `# ` heading is the title, the projected copy becomes
the body, and the `Speaker notes:` bullets from the HTML comment become the
slide's notes. Inline SVG diagrams are noted as a placeholder rather than
rasterised, since that needs a browser; the Marp HTML export renders those.

Usage:
    uv run --no-project python scripts/marp_to_pptx.py 1
    uv run --no-project python scripts/marp_to_pptx.py 1 2 3 4 5
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[1]
DECKS = ROOT / "docs/slides/chicago-sept-24"
OUT = DECKS / "powerpoint"

INK = RGBColor(0x0F, 0x17, 0x2A)
MUT = RGBColor(0x47, 0x55, 0x69)
ACCENT = RGBColor(0x1E, 0x3A, 0x8A)
MONO = "Menlo"


def slides_of(path: Path):
    """Yield (meta, projected_copy, notes) for each slide with a Slide ID."""
    text = path.read_text(encoding="utf-8")
    for seg in text.split("\n---\n"):
        if "Slide ID:" not in seg or "<!--" not in seg:
            continue
        copy, _, rest = seg.partition("<!--")
        note = rest.rpartition("-->")[0]
        meta = {}
        for key in ("Slide ID", "Module", "Instructor", "Type", "Minutes", "Layout"):
            m = re.search(rf"^{re.escape(key)}:\s*(.*)$", note, re.M)
            meta[key] = m.group(1).strip() if m else ""
        bullets = re.findall(r"^- (Say|Ask|Watch|Then):\s*(.*)$", note, re.M)
        src = re.search(r"^Sources:\s*(.*)$", note, re.M)
        yield meta, copy.strip(), (bullets, src.group(1).strip() if src else "")


def clean(line: str) -> str:
    """Strip markdown emphasis and links down to readable text."""
    line = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", line)
    line = line.replace("**", "").replace("`", "")
    return line.strip()


def body_lines(copy: str):
    """Turn projected markdown into (text, level, mono) tuples."""
    out, in_code, in_svg = [], False, False
    for raw in copy.split("\n"):
        s = raw.strip()
        if s.startswith("<svg") or (in_svg and s):
            in_svg = True
            if "</svg>" in s:
                in_svg = False
                out.append(("[diagram — see the HTML deck]", 0, False))
            continue
        if s.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            if s:
                out.append((s, 1, True))
            continue
        if not s or s.startswith(("#", "<")):
            continue
        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if all(set(c) <= set("-: ") for c in cells):
                continue
            out.append((" · ".join(clean(c) for c in cells if c), 1, False))
            continue
        if s.startswith("- "):
            out.append((clean(s[2:]), 1, False))
        elif s.startswith("> "):
            out.append((clean(s[2:]), 0, False))
        else:
            out.append((clean(s), 0, False))
    return out


def build(day: int) -> Path:
    src = DECKS / f"day{day}" / f"day{day}.md"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    count = 0

    for meta, copy, (notes, sources) in slides_of(src):
        slide = prs.slides.add_slide(blank)
        title_txt = ""
        m = re.search(r"^# (.+)$", copy, re.M)
        if m:
            title_txt = clean(m.group(1))

        tb = slide.shapes.add_textbox(Inches(0.62), Inches(0.45), Inches(12.1), Inches(1.0))
        p = tb.text_frame.paragraphs[0]
        p.text = title_txt
        p.font.size = Pt(34)
        p.font.bold = True
        p.font.color.rgb = ACCENT

        rows = body_lines(copy)
        if rows:
            bb = slide.shapes.add_textbox(Inches(0.62), Inches(1.7), Inches(12.1), Inches(4.9))
            tf = bb.text_frame
            tf.word_wrap = True
            for i, (txt, lvl, mono) in enumerate(rows):
                par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                par.text = ("• " + txt) if lvl == 1 and not mono else txt
                par.font.size = Pt(15 if mono else 18)
                par.font.color.rgb = INK if not mono else MUT
                if mono:
                    par.font.name = MONO
                par.space_after = Pt(7)

        ft = slide.shapes.add_textbox(Inches(0.62), Inches(6.85), Inches(12.1), Inches(0.4))
        fp = ft.text_frame.paragraphs[0]
        mins = meta["Minutes"]
        fp.text = f"{meta['Slide ID']}   ·   {mins} min" + (f"   ·   {meta['Layout']}" if meta["Layout"] else "")
        fp.font.size = Pt(10)
        fp.font.color.rgb = MUT

        nt = slide.notes_slide.notes_text_frame
        nt.text = f"{meta['Slide ID']} · {meta['Instructor']} · {mins} min"
        for label, val in notes:
            par = nt.add_paragraph()
            par.text = f"{label}: {clean(val)}"
        if sources:
            par = nt.add_paragraph()
            par.text = f"Sources: {clean(sources)}"
        count += 1

    OUT.mkdir(parents=True, exist_ok=True)
    dest = OUT / f"chicago-sept-24-day{day}.pptx"
    prs.save(dest)
    print(f"day{day}: {count} slides -> {dest.relative_to(ROOT)}")
    return dest


if __name__ == "__main__":
    days = [int(a) for a in sys.argv[1:]] or [1]
    for d in days:
        build(d)
